from __future__ import annotations

import base64
from html import escape
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from weasyprint import HTML


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ASSETS = OUT / "assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

PDF_PATH = OUT / "vaeloryn_product_documentation_v1.pdf"
DOCX_PATH = OUT / "vaeloryn_product_documentation_v1.docx"
HTML_PATH = OUT / "vaeloryn_product_documentation_v1.html"
REPORT_PATH = ROOT / "reports" / "VAELO_Product_Documentation_Final_Report.md"
LOGO_PATH = ROOT / "artifacts" / "vaeloryn" / "public" / "logo.png"

GOLD = "#C99A39"
GOLD_2 = "#E0BA69"
INK = "#111318"
GRAPHITE = "#20242B"
CREAM = "#F5F0E6"
MUTED = "#777D86"
PALE = "#F4F1EA"


def font(size: int, bold: bool = False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def save_flow(path: Path, labels: list[str], title: str):
    width, height = 1800, 720
    im = Image.new("RGB", (width, height), INK)
    draw = ImageDraw.Draw(im)
    draw.text((90, 55), title.upper(), fill=GOLD_2, font=font(34, True))
    left, top, gap = 80, 180, 25
    box_w = (width - 160 - gap * (len(labels) - 1)) // len(labels)
    for i, label in enumerate(labels):
        x = left + i * (box_w + gap)
        draw.rounded_rectangle((x, top, x + box_w, top + 300), radius=22, fill=GRAPHITE, outline=GOLD, width=3)
        words, lines, line = label.split(), [], ""
        for word in words:
            candidate = f"{line} {word}".strip()
            if draw.textlength(candidate, font=font(27, True)) > box_w - 40:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
        total = len(lines) * 42
        for j, text in enumerate(lines):
            tw = draw.textlength(text, font=font(27, True))
            draw.text((x + (box_w - tw) / 2, top + (300 - total) / 2 + j * 42), text, fill=CREAM, font=font(27, True))
        if i < len(labels) - 1:
            ax = x + box_w + 5
            ay = top + 150
            draw.line((ax, ay, ax + gap - 10, ay), fill=GOLD_2, width=5)
            draw.polygon([(ax + gap - 10, ay), (ax + gap - 22, ay - 9), (ax + gap - 22, ay + 9)], fill=GOLD_2)
    draw.text((90, 590), "Conceptual architecture — planned ecosystem; not evidence of operational products or projects.", fill="#BFC3CA", font=font(24))
    im.save(path, quality=95)


def save_allocation(path: Path):
    data = [
        ("Ecosystem & Community", 30),
        ("Public Distribution", 20),
        ("Vaeloryn Treasury", 20),
        ("Team & Contributors", 15),
        ("Founder", 10),
        ("Strategic Partnerships", 5),
    ]
    im = Image.new("RGB", (1800, 880), INK)
    draw = ImageDraw.Draw(im)
    draw.text((90, 55), "CANONICAL VAELO ALLOCATION", fill=GOLD_2, font=font(38, True))
    y = 155
    colors = ["#D6A84B", "#C39031", "#AD7C26", "#896528", "#675127", "#453D2A"]
    for (label, pct), color in zip(data, colors):
        draw.text((90, y + 8), label, fill=CREAM, font=font(28, True))
        draw.rounded_rectangle((570, y, 1650, y + 54), radius=12, fill="#2B3038")
        draw.rounded_rectangle((570, y, 570 + int(1080 * pct / 30), y + 54), radius=12, fill=color)
        draw.text((1670, y + 7), f"{pct}%", fill=CREAM, font=font(28, True))
        y += 105
    draw.text((90, 800), "Fixed total supply: 1,000,000,000 VAELO", fill="#BFC3CA", font=font(28))
    im.save(path, quality=95)


ECOSYSTEM_IMG = ASSETS / "ecosystem_flow.png"
ARCH_IMG = ASSETS / "technical_architecture.png"
ALLOC_IMG = ASSETS / "token_allocation.png"
save_flow(
    ECOSYSTEM_IMG,
    ["Participants", "VAELO ecosystem", "Coordination & participation", "Scientific & technological initiatives", "Development", "Real-world outcomes", "Expanded capacity"],
    "Intended ecosystem flow",
)
save_flow(
    ARCH_IMG,
    ["Deployment factory", "Genesis distribution", "VAELO token", "Founder vesting", "Complete genesis allocation"],
    "Atomic canonical deployment design",
)
save_allocation(ALLOC_IMG)


sections = [
    {
        "n": 1, "title": "Executive overview", "status": "CURRENT STAGE — FOUNDATION / PRE-DEPLOYMENT",
        "body": [
            "Vaeloryn is being developed as a long-term ecosystem intended to connect capital, expertise, technology and opportunity around projects capable of contributing to meaningful advances in science, medicine, energy and frontier technology.",
            "The ambition is to create infrastructure through which resources can be coordinated toward real-world scientific and technological development. The canonical protocol is implemented and tested locally, but the broader ecosystem is not yet fully operational. Base Mainnet deployment, independent security review, the Fjord sale and full ecosystem launch remain incomplete."
        ],
    },
    {
        "n": 2, "title": "The problem", "status": "PROBLEM STATEMENT — NOT A CLAIM OF RESOLUTION",
        "body": ["Promising scientific and technological initiatives can struggle to move from idea to implementation because required resources are fragmented. A coordinated ecosystem could potentially shorten this path by bringing complementary inputs together, while preserving project-specific diligence and execution discipline."],
        "bullets": ["Capital", "Specialist expertise", "Technical and physical infrastructure", "Professional networks", "Development resources", "Long-term support"],
    },
    {
        "n": 3, "title": "The Vaeloryn vision", "status": "LONG-TERM VISION — PLANNED",
        "body": [
            "Vaeloryn’s long-term vision is to help coordinate the inputs required to move credible ideas toward measurable real-world development.",
            "Capital + Expertise + Technology + Execution → Real-world innovation",
            "Intended areas of focus include medical science, energy, artificial intelligence, computing, aerospace, engineering, fundamental science and other frontier technologies. These are areas of interest, not claims that Vaeloryn currently operates projects in every field."
        ],
    },
    {
        "n": 4, "title": "The Vaeloryn ecosystem", "status": "PLANNED ECOSYSTEM ARCHITECTURE",
        "body": [
            "Participants may contribute capital, expertise, ideas, technology or execution capability. The intended VAELO ecosystem would coordinate appropriate forms of participation around selected initiatives.",
            "Initiatives would still require independent evaluation, development, governance, compliance and execution. Real-world outcomes are an objective, not a guaranteed result. Successful work could expand the ecosystem’s future capacity.",
            "Current state: the token protocol, documentation and preparation materials exist; operational products, selected flagship initiatives and a fully functioning participant network remain in development."
        ],
        "image": ECOSYSTEM_IMG,
    },
    {
        "n": 5, "title": "VAELO token", "status": "IMPLEMENTED LOCALLY — NOT DEPLOYED ON BASE MAINNET",
        "body": [
            "VAELO is the fixed-supply token designed for the intended Vaeloryn ecosystem. The target network is Base Mainnet. The canonical deployment package has been prepared, but the canonical token has not been deployed.",
            "Currently implemented functionality: ERC-20 transfers and allowances; ERC20Burnable burn and burnFrom; ERC20Permit / EIP-2612 Permit; EIP-712 typed-data support; and nonce tracking.",
            "The canonical token has no additional mint function, owner or administrator, upgrade mechanism, pause function, blacklist or transfer tax. These design properties do not imply utility, market value or returns."
        ],
        "facts": [("Token", "VAELO"), ("Fixed supply", "1,000,000,000 VAELO"), ("Target network", "Base Mainnet"), ("Standard", "ERC-20"), ("Mainnet status", "NOT DEPLOYED")],
    },
    {
        "n": 6, "title": "Product and token utility", "status": "PLANNED — NOT CURRENT FUNCTIONALITY",
        "body": [
            "Potential future roles may include ecosystem participation, coordination, appropriate access to ecosystem initiatives, support for ecosystem funding mechanisms, participation in future products or services, and interactions among ecosystem participants.",
            "No specific application, staking mechanism, governance system, payment system, revenue-sharing mechanism or current product access is represented as operational. Utility should follow useful products and approved operating models, rather than precede them."
        ],
    },
    {
        "n": 7, "title": "Tokenomics", "status": "FINALIZED CANONICAL ALLOCATION",
        "body": ["The complete fixed supply is 1,000,000,000 VAELO. The six allocations total exactly 100% and are created as one canonical genesis distribution."],
        "table": [
            ["Allocation", "VAELO", "Share"],
            ["Ecosystem & Community", "300,000,000", "30%"],
            ["Public Distribution", "200,000,000", "20%"],
            ["Vaeloryn Treasury", "200,000,000", "20%"],
            ["Team & Contributors", "150,000,000", "15%"],
            ["Founder", "100,000,000", "10%"],
            ["Strategic Partnerships", "50,000,000", "5%"],
            ["TOTAL", "1,000,000,000", "100%"],
        ],
        "image": ALLOC_IMG,
    },
    {
        "n": 8, "title": "Founder vesting", "status": "IMPLEMENTED CANONICAL SCHEDULE",
        "body": ["The founder allocation is 100,000,000 VAELO. Token, beneficiary, start time, linear start and linear end are immutable in the canonical implementation. There is no fourth quarterly release."],
        "table": [
            ["Event", "Timing", "Additional release", "Cumulative"],
            ["Initial", "T0", "2,500,000", "2,500,000"],
            ["Quarter 1", "T0 + 90 days", "2,500,000", "5,000,000"],
            ["Quarter 2", "T0 + 180 days", "2,500,000", "7,500,000"],
            ["Quarter 3", "T0 + 270 days", "2,500,000", "10,000,000"],
            ["Linear phase", "From T0 + 270 days for exactly 1,095 days", "90,000,000 linearly", "100,000,000"],
        ],
    },
    {
        "n": 9, "title": "Technical architecture", "status": "IMPLEMENTED LOCALLY — DEPLOYMENT PENDING",
        "body": [
            "VaelorynDeploymentFactory is a constructor-only, one-shot factory that deploys the distribution, token and founder vesting contracts and completes allocation in the same transaction.",
            "VaelorynGenesisDistribution receives the complete supply directly from the token and performs the one-time six-category allocation. VaelorynToken creates the fixed supply once and includes ERC-20, Burnable and Permit functionality. VaelorynFounderVesting enforces the founder schedule and allocation cap.",
            "The atomic sequence reduces substitution risk between token creation, vesting creation and allocation. It does not mean that a Base Mainnet deployment has occurred."
        ],
        "image": ARCH_IMG,
    },
    {
        "n": 10, "title": "Security and testing", "status": "INTERNAL VERIFICATION COMPLETE — INDEPENDENT AUDIT NOT COMPLETED",
        "body": [
            "The targeted canonical hardening run recorded 27 tests passed, 0 failed and 0 skipped: 25 canonical unit, edge-case, Permit, allocation, vesting, fuzz and factory tests, plus two invariant suites.",
            "Each fuzz target ran 256 times. Each invariant ran 128 times with 8,192 handler calls, for 16,384 invariant handler calls in total.",
            "These are internal development and verification tests and do not constitute an independent security audit. Independent security audit: NOT YET COMPLETED."
        ],
    },
    {
        "n": 11, "title": "Treasury and allocation structure", "status": "APPROVED INTENDED RECIPIENT — NO TRANSFER HAS OCCURRED",
        "body": [
            "The canonical design sends the Vaeloryn Treasury allocation through the one-time genesis distribution to the approved Treasury Safe. The intended use of a Safe supports shared custody and operational controls outside the immutable token contract.",
            "Approved Treasury Safe: 0x3A9bee20360294F8439513Ae70D189f56288baF2",
            "The manifest is marked NOT_DEPLOYED. No canonical VAELO has been transferred to this or any other Mainnet recipient."
        ],
    },
    {
        "n": 12, "title": "Initial public sale", "status": "PROPOSED / WORKING PARAMETERS — SUBJECT TO FINAL FJORD CONFIGURATION",
        "body": [
            "The working concept is a tiered sale on Base using no more than 10,000,000 VAELO, equal to 1% of total supply and carved from the existing Public Distribution allocation. It is not an additional allocation.",
            "The $5,000,000 target reference FDV and $0.005 reference price are planning references, not guaranteed valuations. The sale has not been created, approved, funded or launched."
        ],
        "facts": [("Network", "Base"), ("Sale type", "Tiered"), ("Maximum allocation", "10,000,000 VAELO / 1%"), ("Reference FDV", "$5,000,000"), ("Reference price", "$0.005"), ("Maximum theoretical gross proceeds", "$50,000")],
        "table": [
            ["Tier", "VAELO", "Price", "Gross proceeds"],
            ["1", "1,000,000", "$0.003", "$3,000"],
            ["2", "2,000,000", "$0.004", "$8,000"],
            ["3", "3,000,000", "$0.005", "$15,000"],
            ["4", "4,000,000", "$0.006", "$24,000"],
            ["MAXIMUM", "10,000,000", "Weighted avg. $0.005", "$50,000"],
        ],
    },
    {
        "n": 13, "title": "Use of initial capital", "status": "INTENDED USES — NOT COMPLETED ACTIVITIES",
        "body": ["If an approved sale proceeds, potential uses of initial project funding may include the following. Final budgets, timing and allocations remain subject to approval and available proceeds."],
        "bullets": ["Technical development", "Independent security review", "Infrastructure", "Essential team capability", "Legal and compliance work", "Corporate and operational setup", "Development of the first Vaeloryn initiative", "Communications and transparency infrastructure"],
    },
    {
        "n": 14, "title": "Roadmap", "status": "MILESTONE-GATED — NO DATES PROMISED",
        "body": ["Progression is milestone-gated rather than calendar-gated. Future phases are objectives, not commitments to specific timing or outcomes."],
        "table": [
            ["Phase", "Status", "Representative milestones"],
            ["Phase 1 — Foundation", "IN PROGRESS", "Canonical protocol implemented and tested; website and public documentation established. Community, legal/regulatory guidance, strategy and flagship-project evaluation remain in progress."],
            ["Phase 2 — Execution", "PLANNED", "Build a first flagship project; demonstrate measurable value; grow capability; introduce genuine VAELO utility only where useful."],
            ["Phase 3 — Expansion", "PLANNED", "Expand successful projects; develop additional initiatives; build partnerships; expand the contributor ecosystem; strengthen sustainability."],
            ["Phase 4 — Global Ecosystem", "PLANNED", "Develop a globally connected innovation ecosystem around meaningful science and technology projects."],
        ],
    },
    {
        "n": 15, "title": "Current development status", "status": "STATUS AS OF AUGUST 2026",
        "body": ["The project is at a foundation and preparation stage. The protocol work described below is local implementation and evidence, not a claim of Mainnet launch."],
        "table": [
            ["CURRENT / PREPARED", "NOT YET COMPLETED"],
            ["Canonical contracts implemented", "Base Mainnet deployment"],
            ["Canonical tokenomics finalized", "Mainnet source verification"],
            ["Targeted canonical test suite passing", "Independent security audit"],
            ["Mainnet deployment package prepared", "Fjord sale creation and approval"],
            ["Fjord preparation documentation prepared", "Public token sale"],
            ["Fjord pitch deck prepared", "Full ecosystem launch"],
        ],
    },
    {
        "n": 16, "title": "Transparency and verification", "status": "FUTURE MAINNET VERIFICATION REQUIREMENTS",
        "body": [
            "Canonical identity should be established through the approved Base Mainnet manifest, deployment event and verifiable on-chain evidence—not through a name or placeholder address.",
            "Future publication should include the official VAELO, factory, genesis distribution and founder vesting addresses; deployment transaction; verified source code; compiler and constructor settings; bytecode hashes; tokenomics readback; and vesting readback.",
            "No canonical Mainnet addresses or transaction identifiers exist in the manifest today, so none are inserted here."
        ],
    },
    {
        "n": 17, "title": "Risks and limitations", "status": "MATERIAL UNCERTAINTIES",
        "body": [
            "Vaeloryn is early-stage. Development may be delayed, changed or discontinued. Smart contracts can contain defects despite testing, and an independent audit has not been completed.",
            "Market demand, liquidity and token pricing are uncertain. Regulatory requirements may change or differ across jurisdictions. Adoption depends on participants, useful initiatives and execution capability. Sale operations, eligibility, payment assets, claims and unsold-token treatment require external Fjord confirmation and professional review.",
            "VAELO does not guarantee returns. This document is informational and is not legal, tax, regulatory, financial or investment advice."
        ],
    },
    {
        "n": 18, "title": "Contact and links", "status": "REPOSITORY-VERIFIED LINKS ONLY",
        "body": ["Only links verified in the repository are included. No social URL has been inferred from an account name."],
        "facts": [("Website", "https://vaeloryn.com"), ("GitHub", "https://github.com/Vaeloryn/Vaeloryn"), ("X", "Not supplied in the repository reviewed"), ("Telegram", "Not supplied in the repository reviewed")],
    },
]


def image_uri(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()


def html_table(rows: list[list[str]]) -> str:
    head = "".join(f"<th>{escape(v)}</th>" for v in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{escape(v)}</td>" for v in row) + "</tr>" for row in rows[1:])
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def section_html(s: dict) -> str:
    parts = [f'<section class="section"><div class="kicker">{s["n"]:02d} / PRODUCT DOCUMENTATION</div><h2>{escape(s["title"])}</h2><div class="status">{escape(s["status"])}</div>']
    for p in s.get("body", []):
        cls = "formula" if "→" in p and p.startswith("Capital") else ""
        parts.append(f'<p class="{cls}">{escape(p)}</p>')
    if s.get("facts"):
        parts.append('<div class="facts">' + "".join(f'<div><span>{escape(k)}</span><strong>{escape(v)}</strong></div>' for k, v in s["facts"]) + "</div>")
    if s.get("bullets"):
        parts.append("<ul>" + "".join(f"<li>{escape(x)}</li>" for x in s["bullets"]) + "</ul>")
    if s.get("table"):
        parts.append(html_table(s["table"]))
    if s.get("image"):
        parts.append(f'<img class="diagram" src="{image_uri(s["image"])}" alt="{escape(s["title"])} diagram">')
    parts.append("</section>")
    return "".join(parts)


css = """
@page { size: A4; margin: 18mm 16mm 18mm 16mm; @bottom-right { content: "VAELORYN  •  " counter(page); color: #777D86; font: 8pt Arial; } }
@page:first { margin: 0; @bottom-right { content: none; } }
* { box-sizing: border-box; }
body { margin:0; color:#20242B; font:10.3pt/1.52 Arial, sans-serif; }
.cover { page-break-after:always; height:297mm; background:#090B10; color:#F5F0E6; padding:30mm 22mm; position:relative; overflow:hidden; }
.cover:after { content:""; position:absolute; width:160mm; height:160mm; border:1px solid rgba(201,154,57,.25); border-radius:50%; right:-70mm; bottom:-58mm; }
.cover img { width:30mm; height:30mm; object-fit:contain; }
.cover .brand { letter-spacing:.28em; font-weight:700; font-size:12pt; margin-top:5mm; }
.cover h1 { font:42pt Georgia, serif; letter-spacing:.08em; margin:52mm 0 4mm; }
.cover h3 { color:#E0BA69; font-weight:400; letter-spacing:.1em; margin:0; font-size:16pt; }
.cover .meta { position:absolute; bottom:30mm; left:22mm; right:22mm; border-top:1px solid #8E6B2F; padding-top:7mm; color:#BFC3CA; }
.contents { page-break-after:always; padding-top:8mm; }
.contents h2, .section h2 { font:26pt Georgia, serif; color:#111318; margin:2mm 0 3mm; }
.contents-grid { columns:2; column-gap:12mm; margin-top:10mm; }
.contents-grid div { break-inside:avoid; padding:3mm 0; border-bottom:1px solid #E3DDD0; }
.contents-grid b { color:#C39031; display:inline-block; width:9mm; }
.section { page-break-before:always; }
.kicker { color:#A47626; letter-spacing:.16em; font-size:8pt; font-weight:700; }
.status { display:inline-block; margin:0 0 6mm; padding:2.2mm 3mm; border:1px solid #D8C59D; background:#FBF7EE; color:#77561F; letter-spacing:.06em; font-size:7.8pt; font-weight:700; }
p { margin:0 0 4mm; }
.formula { font:19pt Georgia, serif; color:#8C6828; text-align:center; padding:7mm 3mm; border-top:1px solid #D8C59D; border-bottom:1px solid #D8C59D; }
ul { columns:2; column-gap:10mm; padding-left:5mm; margin:2mm 0 5mm; }
li { margin-bottom:2mm; break-inside:avoid; }
table { width:100%; border-collapse:collapse; margin:5mm 0; font-size:8.6pt; }
th { background:#171A20; color:#F5F0E6; padding:3mm; text-align:left; }
td { border-bottom:1px solid #DDD7CA; padding:2.6mm 3mm; vertical-align:top; }
tbody tr:nth-child(even) td { background:#F7F4EE; }
.facts { display:grid; grid-template-columns:1fr 1fr; gap:2.5mm; margin:4mm 0 6mm; }
.facts div { border-left:3px solid #C99A39; background:#F5F2EB; padding:3mm; }
.facts span { display:block; color:#777D86; text-transform:uppercase; letter-spacing:.08em; font-size:7.5pt; }
.facts strong { display:block; margin-top:1mm; font-size:9.5pt; overflow-wrap:anywhere; }
.diagram { width:100%; margin-top:5mm; border-radius:3mm; }
.note { margin-top:10mm; padding:5mm; background:#F5F2EB; border-left:3px solid #C99A39; }
"""

toc = "".join(f"<div><b>{s['n']:02d}</b>{escape(s['title'])}</div>" for s in sections)
html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Vaeloryn Product Documentation v1.0</title><style>{css}</style></head><body>
<section class="cover"><img src="{image_uri(LOGO_PATH)}"><div class="brand">VAELORYN</div><h1>PRODUCT<br>DOCUMENTATION</h1><h3>VAELO Ecosystem &amp; Product Overview</h3><div class="meta"><b>VERSION 1.0</b><br>August 2026<br><br>Prepared for a Fjord Foundry partnered public-sale application.<br><b>Canonical Base Mainnet deployment: NOT DEPLOYED</b></div></section>
<section class="contents"><div class="kicker">DOCUMENT MAP</div><h2>Contents</h2><p>This product overview distinguishes implemented preparation from planned ecosystem and sale functionality.</p><div class="contents-grid">{toc}</div><div class="note"><b>Important status statement.</b> The canonical VAELO protocol is implemented and internally tested, but has not been deployed on Base Mainnet. No Fjord sale has been created, approved, funded or launched.</div></section>
{''.join(section_html(s) for s in sections)}
</body></html>"""
HTML_PATH.write_text(html, encoding="utf-8")
HTML(filename=str(HTML_PATH), base_url=str(ROOT)).write_pdf(str(PDF_PATH))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("VAELORYN  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(119, 125, 134)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(245, 240, 230)
            if r == 0:
                set_cell_shading(cell, "171A20")
            elif r % 2 == 0:
                set_cell_shading(cell, "F7F4EE")
    doc.add_paragraph()


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9.5)
styles["Title"].font.name = "Georgia"
styles["Title"].font.size = Pt(36)
styles["Heading 1"].font.name = "Georgia"
styles["Heading 1"].font.size = Pt(24)
styles["Heading 1"].font.color.rgb = RGBColor(17, 19, 24)
styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(12)
styles["Heading 2"].font.color.rgb = RGBColor(164, 118, 38)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.15))
p = doc.add_paragraph("V A E L O R Y N", style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("PRODUCT\nDOCUMENTATION", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\nVAELO Ecosystem & Product Overview").font.color.rgb = RGBColor(201, 154, 57)
p = doc.add_paragraph("\nVersion 1.0  •  August 2026\nPrepared for a Fjord Foundry partnered public-sale application.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\nCanonical Base Mainnet deployment: NOT DEPLOYED").bold = True
doc.add_page_break()

doc.add_heading("Contents", level=1)
for s in sections:
    p = doc.add_paragraph()
    r = p.add_run(f"{s['n']:02d}  ")
    r.bold = True
    r.font.color.rgb = RGBColor(201, 154, 57)
    p.add_run(s["title"])
doc.add_paragraph("\nImportant status statement: the canonical VAELO protocol is implemented and internally tested, but has not been deployed on Base Mainnet. No Fjord sale has been created, approved, funded or launched.")

for s in sections:
    doc.add_page_break()
    p = doc.add_paragraph(f"{s['n']:02d} / PRODUCT DOCUMENTATION")
    p.style = "Heading 2"
    doc.add_heading(s["title"], level=1)
    p = doc.add_paragraph()
    r = p.add_run(s["status"])
    r.bold = True
    r.font.color.rgb = RGBColor(119, 86, 31)
    set_cell = None
    for text in s.get("body", []):
        p = doc.add_paragraph(text)
        if "→" in text and text.startswith("Capital"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Georgia"
                run.font.size = Pt(17)
                run.font.color.rgb = RGBColor(140, 104, 40)
    if s.get("facts"):
        add_table(doc, [["Item", "Verified or stated value"]] + [[k, v] for k, v in s["facts"]])
    for item in s.get("bullets", []):
        doc.add_paragraph(item, style="List Bullet")
    if s.get("table"):
        add_table(doc, s["table"])
    if s.get("image"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(s["image"]), width=Inches(6.95))

for section in doc.sections:
    add_page_number(section.footer.paragraphs[0])
doc.core_properties.title = "Vaeloryn Product Documentation"
doc.core_properties.subject = "VAELO Ecosystem & Product Overview"
doc.core_properties.author = "Vaeloryn"
doc.core_properties.comments = "Canonical Base Mainnet deployment not completed."
doc.save(DOCX_PATH)

REPORT_PATH.write_text(
    """# VAELO Product Documentation — Final Report

## Files created
- `deliverables/vaeloryn_product_documentation_v1.pdf`
- `deliverables/vaeloryn_product_documentation_v1.docx`
- `deliverables/vaeloryn_product_documentation_v1.html` (shared styled generation source)

## Sections included
The PDF contains 20 A4 pages: a cover, contents page and all 18 requested sections. Those sections are executive overview; problem; vision; ecosystem; token; planned utility; tokenomics; founder vesting; technical architecture; security and testing; treasury and allocations; proposed sale; intended capital use; roadmap; current status; transparency and verification; risks; and contact/links.

## Primary sources used
- Attached product-documentation brief
- Canonical Solidity contracts in `vaelo_canonical/src/`
- Base Mainnet manifest and Fjord sale specification
- Canonical implementation, testing and founder-vesting reports
- Vaeloryn roadmap and repository metadata

## Missing information
- No verified X URL was found; the document marks it as not supplied.
- No verified Telegram URL was found; the document marks it as not supplied.
- Canonical Base Mainnet addresses, transaction data and verification URLs remain unavailable because deployment has not occurred.
- Final Fjord configuration, timing, accepted assets, eligibility, claims, fees and approval remain externally unconfirmed.

## Claims intentionally excluded
No claims of active partnerships, users, revenue, investors, independent audit completion, Base Mainnet deployment, Fjord approval, live sale, guaranteed valuation, operational utility, returns, equity, security status or ownership rights were added.

## Action confirmation
No deployment, transaction broadcast, VAELO/ETH/USDC transfer, sale creation, canonical Solidity modification or website modification was performed.

**FINAL STATUS: PRODUCT DOCUMENTATION CREATED — NO DEPLOYMENT OR SALE ACTION PERFORMED**
""",
    encoding="utf-8",
)

print(PDF_PATH)
print(DOCX_PATH)
print(REPORT_PATH)